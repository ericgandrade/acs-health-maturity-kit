#!/usr/bin/env python3
"""Build toolkit/ACS-Health-Maturity-Assessment.xlsx from the CSV files in data/.

Usage:  python tools/build_workbook.py [output.xlsx]
The CSV files are the single source of truth; edit them and rebuild.
"""
import csv, math, sys
from collections import OrderedDict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.chart import BarChart, RadarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

VERSION = '0.2.0'
RELEASE_DATE = '2026-10-29'
ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'toolkit' / 'ACS-Health-Maturity-Assessment.xlsx'

BLUE, DARK, GREY, INPUT, LIGHT = '0090C7', '262626', 'D1D2D4', 'FFF2CC', 'EAF5FB'
F = 'Arial'
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)

SERVICES = OrderedDict([
    ('License Optimization', 'License Optimization'),
    ('App Modernization', 'App Modernization'),
    ('Zero Trust', 'Zero Trust'),
])
ANSWERS = ['0 - Not in place', '1 - Initial', '2 - Developing', '3 - Defined', '4 - Optimized', 'N/A']
LEVELS = [('0 - Not in place', 'The practice does not exist.'),
          ('1 - Initial', 'Ad hoc; depends on individuals; not documented.'),
          ('2 - Developing', 'Partly implemented or documented; not applied consistently.'),
          ('3 - Defined', 'Documented and applied consistently across the organization.'),
          ('4 - Optimized', 'Measured, reviewed regularly and improved; automated where practical.')]
CPGS = [('Essential', 'Mitigate Known Vulnerabilities'), ('Essential', 'Email Security'),
        ('Essential', 'Multifactor Authentication'), ('Essential', 'Basic Cybersecurity Training'),
        ('Essential', 'Strong Encryption'), ('Essential', 'Revoke Credentials for Departing Workforce Members'),
        ('Essential', 'Basic Incident Planning and Preparedness'), ('Essential', 'Unique Credentials'),
        ('Essential', 'Separate User and Privileged Accounts'), ('Essential', 'Vendor/Supplier Cybersecurity Requirements'),
        ('Enhanced', 'Asset Inventory'), ('Enhanced', 'Third Party Vulnerability Disclosure'),
        ('Enhanced', 'Third Party Incident Reporting'), ('Enhanced', 'Cybersecurity Testing'),
        ('Enhanced', 'Cybersecurity Mitigation'), ('Enhanced', 'Detect and Respond to Relevant Threats and TTPs'),
        ('Enhanced', 'Network Segmentation'), ('Enhanced', 'Centralized Log Collection'),
        ('Enhanced', 'Centralized Incident Planning and Preparedness'), ('Enhanced', 'Configuration Management')]


def font(**kw):
    base = dict(name=F, size=10, color=DARK)
    base.update(kw)
    return Font(**base)


def header(ws, row, labels, widths=None):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font = font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.alignment = CENTER
        cell.border = BORDER
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = font(size=16, bold=True, color=BLUE)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = font(size=10, italic=True)
    ws.sheet_view.showGridLines = False


def fit_height(ws, row, texts_widths, base=13.0):
    lines = 1
    for text, width in texts_widths:
        if text:
            n = sum(max(1, math.ceil(len(part) / max(width * 1.15, 1))) for part in str(text).split('\n'))
            lines = max(lines, n)
    ws.row_dimensions[row].height = min(409.0, max(base, 13.0 * lines + 2))


def split_list(text):
    """List fields in the CSV are separated by ' | '."""
    return [t.strip() for t in (text or '').split(' | ') if t.strip()]


def bullets(text):
    return '\n'.join('• ' + t for t in split_list(text))


def level_guide(q):
    texts = [q['level_0'], q['level_1'], q['level_2'], q['level_3'], q['target_state']]
    return '\n'.join(f'{lvl[0]} – {t}' for lvl, t in zip(LEVELS, texts))


def load_csv(name):
    with open(ROOT / 'data' / name, encoding='utf-8') as f:
        return list(csv.DictReader(f))


def build():
    questions = load_csv('questionnaire.csv')
    lic_map = load_csv('license_security_map.csv')
    wb = Workbook()

    # ---------------- Lists (hidden) ----------------
    ls = wb.active
    ls.title = 'Lists'
    lists = {
        'A': ('Answers', ANSWERS),
        'B': ('Levels', [l[0][4:] for l in LEVELS]),
        'C': ('YesNoUnknown', ['Yes', 'No', 'Unknown']),
        'D': ('YesNo', ['Yes', 'No']),
        'E': ('Hosting', ['On-premises', 'Cloud IaaS', 'Cloud PaaS', 'SaaS', 'Hybrid']),
        'F': ('Decision', ['Retire', 'Retain', 'Rehost', 'Replatform', 'Refactor', 'Replace', 'Isolate']),
        'G': ('OrgType', ['Critical access hospital', 'Rural emergency hospital', 'Community hospital', 'Health system',
                          'Academic medical center', 'Clinic or physician group', 'Other']),
        'H': ('Agreement', ['Enterprise Agreement', 'Cloud Solution Provider (CSP)', 'Microsoft Customer Agreement (MCA)', 'Other or unknown']),
        'I': ('M365Plan', ['Microsoft 365 E3', 'Microsoft 365 E5', 'Microsoft 365 Business Premium', 'Mixed', 'Other or unknown']),
        'J': ('Deployed', ['Yes', 'Partial', 'No']),
        'K': ('OneToFive', ['1', '2', '3', '4', '5']),
    }
    for col, (name, vals) in lists.items():
        ls[f'{col}1'] = name
        for i, v in enumerate(vals, 2):
            ls[f'{col}{i}'] = int(v) if name == 'OneToFive' else v
    ls.sheet_state = 'hidden'

    def dv(ws, col, n, rng):
        v = DataValidation(type='list', formula1=f'=Lists!${col}$2:${col}${n + 1}', allow_blank=True)
        v.error = 'Choose a value from the list.'
        v.errorTitle = 'Invalid entry'
        ws.add_data_validation(v)
        v.add(rng)

    # ---------------- Start Here ----------------
    st = wb.create_sheet('Start Here', 0)
    title(st, 'ACS Health Maturity Kit — Assessment Workbook',
          f'Version {VERSION} · {RELEASE_DATE} · Andrade Consulting Services · Licensed under CC BY 4.0')
    st.column_dimensions['A'].width = 26
    st.column_dimensions['B'].width = 95
    rows = [
        ('Purpose', 'A free self-assessment for U.S. healthcare organizations of their maturity in three areas: '
                    'License Optimization, App Modernization and Zero Trust. Questions are mapped to the HHS Healthcare and '
                    'Public Health Cybersecurity Performance Goals (CPGs), the HIPAA Security Rule, NIST CSF 2.0, the CISA Zero '
                    'Trust Maturity Model v2.0, NIST SP 800-218 (SSDF), ISO/IEC 19770-1 and the FinOps Framework.'),
        ('Who should answer', 'The IT leader, with input from security, infrastructure, applications, procurement and '
                              'clinical informatics. Answer based on evidence, not intent.'),
        ('How to use', '1) Fill in the Profile sheet. 2) Answer every question in the three service sheets using the '
                       'yellow Answer column: pick the level in the "Maturity levels (0–4)" column that matches what the '
                       'evidence shows today. 3) Use the Question Guide sheet for interview probes, evidence, red flags and '
                       'the path to the next level. 4) Review the Dashboard, Priority Gaps and CPG Crosswalk sheets. '
                       '5) Use App Portfolio Triage and License-Security Map for detailed planning.'),
        ('Rule of evidence', 'Choose the highest level whose description is fully true today. If part of it is missing, '
                             'choose the level below. Plans and pilots do not count. Scores of 3 or 4 require evidence you have seen.'),
        ('Cells to edit', 'Only yellow cells. All other values are calculated.'),
        ('Example', 'ZT-01: if MFA protects the VPN but not email or administrator accounts, the level 1 description is fully true and '
                    'level 2 is not, so answer "1 - Initial". If MFA covers remote access and email for most users but '
                    'vendor accounts are excluded, answer "2 - Developing".'),
        ('Weights', '3 = critical to patient safety or data protection; 2 = important; 1 = foundational.'),
        ('Scoring', 'Each answer scores 0–4. Scores are weighted averages; "N/A" and blank answers are excluded. '
                    'Gap points = (4 − score) × weight and rank the Priority Gaps.'),
        ('Privacy', 'Do NOT enter protected health information (PHI), passwords or network addresses in this workbook.'),
        ('Disclaimer', 'This is a self-assessment aid. It is not legal, licensing or compliance advice and does not certify '
                       'compliance with HIPAA or any other requirement.'),
        ('Report use and feedback', 'Please tell us how you used the kit: open an "Adoption report" issue in the project '
                                    'repository. Usage reports help improve the kit for other organizations.'),
    ]
    r = 4
    for k, v in rows:
        st.cell(row=r, column=1, value=k).font = font(bold=True)
        c = st.cell(row=r, column=2, value=v)
        c.font = font()
        c.alignment = WRAP
        st.cell(row=r, column=1).alignment = WRAP
        fit_height(st, r, [(v, 95)])
        r += 1
    r += 1
    st.cell(row=r, column=1, value='Maturity scale').font = font(bold=True, color=BLUE, size=12)
    r += 1
    header(st, r, ['Answer', 'Meaning'])
    for lvl, meaning in LEVELS:
        r += 1
        st.cell(row=r, column=1, value=lvl).font = font(bold=True)
        st.cell(row=r, column=2, value=meaning).font = font()
        for c in (1, 2):
            st.cell(row=r, column=c).border = BORDER
    r += 2
    st.cell(row=r, column=1, value='Sheets').font = font(bold=True, color=BLUE, size=12)
    sheets_desc = [
        ('Profile', 'Organization context.'),
        ('License Optimization', f'{sum(q["service"]=="License Optimization" for q in questions)} questions on licensing governance, entitlements, usage, security value, cloud cost and programs.'),
        ('App Modernization', f'{sum(q["service"]=="App Modernization" for q in questions)} questions on portfolio visibility, lifecycle, roadmap, engineering, resilience and cloud readiness.'),
        ('Zero Trust', f'{sum(q["service"]=="Zero Trust" for q in questions)} questions organized by the CISA Zero Trust Maturity Model pillars plus governance and readiness.'),
        ('Dashboard', 'Scores by service and dimension, with charts.'),
        ('Priority Gaps', 'The 20 gaps with the most gap points and the recommended action for each.'),
        ('CPG Crosswalk', 'Status against each of the 20 HHS Cybersecurity Performance Goals.'),
        ('App Portfolio Triage', 'Application inventory with risk scoring and a suggested disposition.'),
        ('License-Security Map', 'Security capabilities included in common Microsoft licenses and whether they are deployed.'),
        ('Question Guide', 'For every question: interview probes, evidence and how to obtain it, red flags, healthcare considerations, the path to the next level and sources.'),
    ]
    for name, desc in sheets_desc:
        r += 1
        st.cell(row=r, column=1, value=name).font = font(bold=True)
        st.cell(row=r, column=2, value=desc).font = font()
    r += 2
    st.cell(row=r, column=1, value='Legend').font = font(bold=True, color=BLUE, size=12)
    r += 1
    st.cell(row=r, column=1, value='Input cell').fill = PatternFill('solid', fgColor=INPUT)
    st.cell(row=r, column=1).border = BORDER
    st.cell(row=r, column=2, value='Yellow cells are for your answers.').font = font()

    # ---------------- Profile ----------------
    pf = wb.create_sheet('Profile', 1)
    title(pf, 'Organization Profile', 'Context for interpreting the results. Do not enter PHI.')
    pf.column_dimensions['A'].width = 34
    pf.column_dimensions['B'].width = 44
    prof = [('Organization name', None), ('State', None), ('Organization type', 'G7'),
            ('Licensed beds', None), ('Workforce size (approx.)', None), ('IT staff (FTE)', None),
            ('Primary EHR', None), ('Microsoft agreement', 'H4'), ('Microsoft 365 plan', 'I5'),
            ('Assessment date', None), ('Assessed by (role)', None)]
    for i, (lab, lst) in enumerate(prof, 4):
        pf.cell(row=i, column=1, value=lab).font = font(bold=True)
        c = pf.cell(row=i, column=2)
        c.fill = PatternFill('solid', fgColor=INPUT)
        c.border = BORDER
        c.font = font()
        if lst:
            dv(pf, lst[0], int(lst[1:]), f'B{i}')
    pf['B13'].number_format = 'yyyy-mm-dd'

    # ---------------- Questionnaire sheets ----------------
    QH = ['ID', 'Dimension', 'Question', 'Maturity levels (0–4)', 'Evidence to review',
          'HHS CPG mapping', 'Other references', 'Weight', 'Answer', 'Notes', 'Recommended action if gap',
          'Score', 'Weighted score', 'Weight applied', 'Gap points']
    QW = [8, 18, 40, 60, 42, 26, 26, 8, 17, 26, 36, 8, 9, 9, 9]
    FIRST = 5
    service_rows = {}
    for idx, svc in enumerate(SERVICES, 2):
        ws = wb.create_sheet(svc, idx)
        title(ws, f'{svc} — Maturity Questionnaire',
              'Answer each question in the yellow column using the 0–4 scale (see Start Here). Grey columns are calculated.')
        header(ws, FIRST - 1, QH, QW)
        ws.freeze_panes = ws.cell(row=FIRST, column=4)
        items = [q for q in questions if q['service'] == svc]
        r = FIRST
        for q in items:
            vals = [q['id'], q['dimension'], q['question'], level_guide(q), bullets(q['evidence']), q['hhs_cpg'],
                    q['references'], int(q['weight']), None, None, q['recommended_action']]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = font(bold=(c == 1))
                cell.alignment = CENTER if c in (1, 8) else WRAP
                cell.border = BORDER
            for c in (9, 10):
                ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=INPUT)
            ws.cell(row=r, column=12, value=f'=IF(OR(I{r}="",I{r}="N/A"),"",VALUE(LEFT(I{r},1)))')
            ws.cell(row=r, column=13, value=f'=IF(L{r}="",0,L{r}*H{r})')
            ws.cell(row=r, column=14, value=f'=IF(L{r}="",0,H{r})')
            ws.cell(row=r, column=15, value=f'=IF(L{r}="",0,(4-L{r})*H{r})')
            for c in range(12, 16):
                cell = ws.cell(row=r, column=c)
                cell.font = font(color='595959')
                cell.fill = PatternFill('solid', fgColor='F2F2F2')
                cell.alignment = CENTER
                cell.border = BORDER
            fit_height(ws, r, [(vals[2], 40), (vals[3], 60), (vals[4], 42), (vals[10], 36), (vals[5], 26), (vals[6], 26)])
            r += 1
        last = r - 1
        dv(ws, 'A', len(ANSWERS), f'I{FIRST}:I{last}')
        ws.conditional_formatting.add(f'L{FIRST}:L{last}', CellIsRule(operator='between', formula=['0', '1'], fill=PatternFill('solid', fgColor='F8CBAD')))
        ws.conditional_formatting.add(f'L{FIRST}:L{last}', CellIsRule(operator='equal', formula=['2'], fill=PatternFill('solid', fgColor='FFE699')))
        ws.conditional_formatting.add(f'L{FIRST}:L{last}', CellIsRule(operator='between', formula=['3', '4'], fill=PatternFill('solid', fgColor='C6EFCE')))
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f'{FIRST - 1}:{FIRST - 1}'
        service_rows[svc] = (FIRST, last, items)

    # ---------------- AllItems (hidden helper) ----------------
    al = wb.create_sheet('AllItems')
    al.append(['Service', 'ID', 'Dimension', 'Question', 'Recommended action', 'HHS CPG', 'Score',
               'Weight applied', 'Weighted score', 'Gap points', 'Rank key', 'Weight'])
    ar = 2
    for svc, (first, last, items) in service_rows.items():
        sref = f"'{svc}'"
        for i in range(first, last + 1):
            al.cell(row=ar, column=1, value=svc)
            for c, src in ((2, 'A'), (3, 'B'), (4, 'C'), (5, 'K'), (6, 'F')):
                al.cell(row=ar, column=c, value=f'={sref}!{src}{i}&""')
            al.cell(row=ar, column=7, value=f'={sref}!L{i}')
            al.cell(row=ar, column=8, value=f'={sref}!N{i}')
            al.cell(row=ar, column=9, value=f'={sref}!M{i}')
            al.cell(row=ar, column=10, value=f'={sref}!O{i}')
            al.cell(row=ar, column=11, value=f'=J{ar}+(1000-ROW())/100000')
            al.cell(row=ar, column=12, value=f'={sref}!H{i}')
            ar += 1
    alast = ar - 1
    al.sheet_state = 'hidden'
    A = lambda col: f'AllItems!${col}$2:${col}${alast}'

    # ---------------- Dashboard ----------------
    db = wb.create_sheet('Dashboard', 5)
    title(db, 'Dashboard', 'Weighted maturity scores on a 0–4 scale. Unanswered and N/A questions are excluded.')
    for col, w in zip('ABCDEFG', [30, 14, 12, 20, 12, 12, 4]):
        db.column_dimensions[col].width = w
    header(db, 4, ['Service', 'Score (0–4)', 'Score %', 'Maturity level', 'Answered', 'Questions'])
    r = 5
    for svc in SERVICES:
        db.cell(row=r, column=1, value=svc)
        db.cell(row=r, column=2, value=f'=IFERROR(SUMIFS({A("I")},{A("A")},A{r})/SUMIFS({A("H")},{A("A")},A{r}),"")')
        db.cell(row=r, column=6, value=f'=COUNTIFS({A("A")},A{r})')
        r += 1
    db.cell(row=r, column=1, value='Overall')
    db.cell(row=r, column=2, value=f'=IFERROR(SUM({A("I")})/SUM({A("H")}),"")')
    db.cell(row=r, column=6, value=f'=SUM(F5:F{r - 1})')
    for rr in range(5, r + 1):
        db.cell(row=rr, column=3, value=f'=IF(B{rr}="","",B{rr}/4)')
        db.cell(row=rr, column=4, value=f'=IF(B{rr}="","Not assessed",INDEX(Lists!$B$2:$B$6,MIN(5,INT(B{rr})+1)))')
        if rr < r:
            db.cell(row=rr, column=5, value=f'=COUNTIFS({A("A")},A{rr},{A("G")},">=0")')
        else:
            db.cell(row=rr, column=5, value=f'=SUM(E5:E{r - 1})')
        for c in range(1, 7):
            cell = db.cell(row=rr, column=c)
            cell.border = BORDER
            cell.font = font(bold=(rr == r or c == 1))
        db.cell(row=rr, column=2).number_format = '0.00'
        db.cell(row=rr, column=3).number_format = '0%'
    overall_row = r

    # HHS CPG summary
    r = overall_row + 2
    db.cell(row=r, column=1, value='HHS Cybersecurity Performance Goals').font = font(bold=True, color=BLUE, size=12)
    r += 1
    header(db, r, ['Goal type', 'Goals', 'Meets', 'Partially meets', 'Gap', 'Not assessed'])
    for i, t in enumerate(['Essential', 'Enhanced']):
        rr = r + 1 + i
        db.cell(row=rr, column=1, value=t)
        db.cell(row=rr, column=2, value=f'=COUNTIFS(\'CPG Crosswalk\'!$A$5:$A$24,A{rr})')
        for c, status in zip((3, 4, 5, 6), ('Meets*', 'Partially*', 'Gap', 'Not assessed')):
            db.cell(row=rr, column=c, value=f'=COUNTIFS(\'CPG Crosswalk\'!$A$5:$A$24,A{rr},\'CPG Crosswalk\'!$E$5:$E$24,"{status}")')
        for c in range(1, 7):
            db.cell(row=rr, column=c).border = BORDER
            db.cell(row=rr, column=c).font = font(bold=(c == 1))
    r = r + 4

    # Dimension tables + radar charts
    dim_blocks = []
    for svc, (first, last, items) in service_rows.items():
        dims = list(OrderedDict.fromkeys(q['dimension'] for q in items))
        db.cell(row=r, column=1, value=f'{svc} — by dimension').font = font(bold=True, color=BLUE, size=12)
        r += 1
        header(db, r, ['Dimension', 'Score (0–4)', 'Score %', 'Maturity level'])
        start = r + 1
        for d in dims:
            r += 1
            db.cell(row=r, column=1, value=d)
            db.cell(row=r, column=2, value=f'=IFERROR(SUMIFS({A("I")},{A("A")},"{svc}",{A("C")},A{r})/SUMIFS({A("H")},{A("A")},"{svc}",{A("C")},A{r}),"")')
            db.cell(row=r, column=3, value=f'=IF(B{r}="","",B{r}/4)')
            db.cell(row=r, column=4, value=f'=IF(B{r}="","Not assessed",INDEX(Lists!$B$2:$B$6,MIN(5,INT(B{r})+1)))')
            db.cell(row=r, column=2).number_format = '0.00'
            db.cell(row=r, column=3).number_format = '0%'
            for c in range(1, 5):
                db.cell(row=r, column=c).border = BORDER
                db.cell(row=r, column=c).font = font()
        dim_blocks.append((svc, start, r))
        r += 3

    bar = BarChart()
    bar.type = 'bar'
    bar.title = 'Maturity by service (0–4)'
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 4
    bar.y_axis.majorUnit = 1
    bar.add_data(Reference(db, min_col=2, min_row=4, max_row=overall_row), titles_from_data=True)
    bar.set_categories(Reference(db, min_col=1, min_row=5, max_row=overall_row))
    bar.legend = None
    bar.height, bar.width = 7, 14
    db.add_chart(bar, 'H4')
    anchor_row = 20
    for svc, s, e in dim_blocks:
        rc = RadarChart()
        rc.type = 'marker'
        rc.style = 26
        rc.title = f'{svc} by dimension'
        rc.add_data(Reference(db, min_col=2, min_row=s - 1, max_row=e), titles_from_data=True)
        rc.set_categories(Reference(db, min_col=1, min_row=s, max_row=e))
        rc.y_axis.scaling.min = 0
        rc.y_axis.scaling.max = 4
        rc.y_axis.majorUnit = 1
        rc.legend = None
        rc.height, rc.width = 8, 14
        db.add_chart(rc, f'H{anchor_row}')
        anchor_row += 17

    # ---------------- Priority Gaps ----------------
    pg = wb.create_sheet('Priority Gaps', 6)
    title(pg, 'Priority Gaps', 'The 20 answered items with the most gap points ((4 − score) × weight). Updates automatically.')
    header(pg, 4, ['Rank', 'Service', 'ID', 'Dimension', 'Question', 'Score', 'Weight', 'Gap points',
                   'Recommended action', 'HHS CPG mapping', 'Key', 'Row'],
           [6, 18, 8, 20, 48, 7, 7, 9, 44, 30, 8, 6])
    for n in range(1, 21):
        r = 4 + n
        pg.cell(row=r, column=1, value=n)
        pg.cell(row=r, column=11, value=f'=LARGE({A("K")},A{r})')
        pg.cell(row=r, column=12, value=f'=MATCH(K{r},{A("K")},0)')
        g = f'INDEX({A("J")},L{r})'
        for c, col in ((2, 'A'), (3, 'B'), (4, 'C'), (5, 'D'), (6, 'G'), (7, 'L'), (8, 'J'), (9, 'E'), (10, 'F')):
            pg.cell(row=r, column=c, value=f'=IF({g}>0,INDEX({A(col)},L{r}),"")')
        for c in range(1, 13):
            cell = pg.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = font(color='A6A6A6' if c > 10 else DARK)
            cell.alignment = CENTER if c in (1, 3, 6, 7, 8) else WRAP
        pg.row_dimensions[r].height = 40
    pg.column_dimensions['K'].hidden = True
    pg.column_dimensions['L'].hidden = True
    pg.freeze_panes = 'B5'
    pg['A26'] = 'Copy final priorities into templates/remediation-roadmap.md and assign an owner and target date to each.'
    pg['A26'].font = font(italic=True)

    # ---------------- CPG Crosswalk ----------------
    cw = wb.create_sheet('CPG Crosswalk', 7)
    title(cw, 'HHS Cybersecurity Performance Goals — Crosswalk',
          'Average score of the questions mapped to each goal. Meets = average of 3 or more.')
    header(cw, 4, ['Goal type', 'Goal', 'Related questions', 'Average score', 'Status', 'Tag'],
           [12, 52, 34, 14, 26, 10])
    for i, (t, g) in enumerate(CPGS):
        r = 5 + i
        tag = f'CPG-{"E" if t == "Essential" else "X"}: {g}'
        related = ', '.join(q['id'] for q in questions if tag in q['hhs_cpg'])
        cw.cell(row=r, column=1, value=t)
        cw.cell(row=r, column=2, value=g)
        cw.cell(row=r, column=3, value=related)
        cw.cell(row=r, column=6, value=tag)
        cw.cell(row=r, column=4, value=f'=IFERROR(AVERAGEIFS({A("G")},{A("F")},"*"&F{r}&"*"),"")')
        cw.cell(row=r, column=5, value=f'=IF(D{r}="","Not assessed",IF(D{r}>=3,"Meets (Defined or above)",IF(D{r}>=2,"Partially meets","Gap")))')
        cw.cell(row=r, column=4).number_format = '0.00'
        for c in range(1, 7):
            cell = cw.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = font(color='A6A6A6' if c == 6 else DARK)
            cell.alignment = Alignment(vertical='top') if c == 6 else WRAP
        cw.row_dimensions[r].height = 28
    cw.column_dimensions['F'].hidden = True
    cw.conditional_formatting.add('E5:E24', FormulaRule(formula=['LEFT(E5,5)="Meets"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    cw.conditional_formatting.add('E5:E24', FormulaRule(formula=['LEFT(E5,9)="Partially"'], fill=PatternFill('solid', fgColor='FFE699')))
    cw.conditional_formatting.add('E5:E24', FormulaRule(formula=['E5="Gap"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    cw['A26'] = 'Source: U.S. Department of Health and Human Services, Healthcare and Public Health Sector Cybersecurity Performance Goals, https://hhscyber.hhs.gov/cybersecurity-performance-goals.html'
    cw['A26'].font = font(italic=True, size=9)

    # ---------------- App Portfolio Triage ----------------
    tp = wb.create_sheet('App Portfolio Triage', 8)
    title(tp, 'App Portfolio Triage', 'One row per application. Yellow = input. Technical risk, priority and suggested disposition are calculated.')
    tp['A3'] = ('Technical risk (0–8): unsupported platform 3 (unknown 2) + open critical vulnerabilities 3 (unknown 2) + PHI 1 + more than 5 integrations 1. '
                'Priority = technical risk × criticality. The suggested disposition is a starting point; owners make the final decision.')
    tp['A3'].font = font(italic=True, size=9)
    TH = ['Application', 'Business function', 'Vendor', 'Hosting', 'Platform / OS', 'Platform supported?',
          'Open critical vulnerabilities?', 'Handles PHI?', 'Integrations (count)', 'Criticality (1–5)', 'Business value (1–5)',
          'Annual cost (USD)', 'Technical risk (0–8)', 'Priority (0–40)', 'Suggested disposition', 'Final decision', 'Owner', 'Target date', 'Notes']
    header(tp, 5, TH, [24, 20, 16, 13, 16, 12, 14, 10, 12, 11, 11, 13, 11, 10, 30, 13, 16, 12, 24])
    examples = [
        ['EXAMPLE – Laboratory information system', 'Laboratory results', 'Vendor A', 'On-premises', 'Windows Server 2012 R2',
         'No', 'Unknown', 'Yes', 7, 5, 4, 180000],
        ['EXAMPLE – Staff scheduling tool', 'Workforce scheduling', 'Vendor B', 'SaaS', 'Vendor-managed',
         'Yes', 'No', 'No', 1, 2, 2, 12000],
    ]
    for r in range(6, 6 + 52):
        ex = examples[r - 6] if r - 6 < len(examples) else None
        for c in range(1, 20):
            cell = tp.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = font(italic=bool(ex), color='7F7F7F' if ex else DARK)
            cell.alignment = WRAP
            if c <= 12 or c >= 16:
                cell.fill = PatternFill('solid', fgColor=INPUT)
            if ex and c <= 12:
                cell.value = ex[c - 1]
        tp.cell(row=r, column=13, value=f'=IF(A{r}="","",IF(F{r}="No",3,IF(F{r}="Unknown",2,0))+IF(G{r}="Yes",3,IF(G{r}="Unknown",2,0))+IF(H{r}="Yes",1,0)+IF(N(I{r})>5,1,0))')
        tp.cell(row=r, column=14, value=f'=IF(OR(M{r}="",J{r}=""),"",M{r}*J{r})')
        tp.cell(row=r, column=15, value=(f'=IF(OR(M{r}="",J{r}="",K{r}=""),"",IF(AND(K{r}<=2,J{r}<=2),"Retire or consolidate",'
                                          f'IF(AND(M{r}>=5,K{r}>=3),"Modernize first (upgrade or replatform)",'
                                          f'IF(M{r}>=5,"Isolate and plan replacement",IF(K{r}>=4,"Retain and invest","Retain and monitor")))))'))
        tp.cell(row=r, column=12).number_format = '$#,##0'
        tp.cell(row=r, column=18).number_format = 'yyyy-mm-dd'
        for c in (13, 14, 15):
            tp.cell(row=r, column=c).fill = PatternFill('solid', fgColor='F2F2F2')
    tlast = 6 + 51
    dv(tp, 'E', 5, f'D6:D{tlast}')
    dv(tp, 'C', 3, f'F6:G{tlast}')
    dv(tp, 'D', 2, f'H6:H{tlast}')
    dv(tp, 'K', 5, f'J6:K{tlast}')
    dv(tp, 'F', 7, f'P6:P{tlast}')
    tp.freeze_panes = 'B6'
    tp.conditional_formatting.add(f'N6:N{tlast}', CellIsRule(operator='greaterThanOrEqual', formula=['20'], fill=PatternFill('solid', fgColor='F8CBAD')))

    # ---------------- License-Security Map ----------------
    lm = wb.create_sheet('License-Security Map', 9)
    title(lm, 'License-Security Map', 'Security capabilities included in common Microsoft licenses, and whether you have deployed them.')
    lm['A3'] = ('Licensing contents and names change frequently. Verify each line against current Microsoft documentation and your agreement '
                'before acting. This sheet is not licensing or legal advice.')
    lm['A3'].font = font(italic=True, size=9, color='C00000')
    header(lm, 5, ['Security capability', 'Typical Microsoft licensing (verify)', 'HHS CPG supported', 'Related questions',
                   'Licensed?', 'Deployed?', 'Status', 'Notes'], [40, 48, 34, 16, 11, 11, 34, 26])
    for i, row in enumerate(lic_map):
        r = 6 + i
        vals = [row['capability'], row['typical_microsoft_licensing'], row['hhs_cpg'].replace('CPG-E: ', 'Essential: ').replace('CPG-X: ', 'Enhanced: '), row['related_questions']]
        for c, v in enumerate(vals, 1):
            lm.cell(row=r, column=c, value=v)
        lm.cell(row=r, column=7, value=(f'=IF(E{r}="","",IF(E{r}="Yes",IF(F{r}="Yes","Deployed",IF(F{r}="","Check deployment",'
                                          f'"Licensed but not fully deployed – quick win")),IF(E{r}="No","Not licensed – evaluate need","Confirm licensing")))'))
        for c in range(1, 9):
            cell = lm.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = font()
            cell.alignment = WRAP
            if c in (5, 6, 8):
                cell.fill = PatternFill('solid', fgColor=INPUT)
        fit_height(lm, r, [(vals[0], 40), (vals[1], 48), (vals[2], 34)])
    llast = 5 + len(lic_map)
    dv(lm, 'C', 3, f'E6:E{llast}')
    dv(lm, 'J', 3, f'F6:F{llast}')
    lm.conditional_formatting.add(f'G6:G{llast}', FormulaRule(formula=['LEFT(G6,8)="Licensed"'], fill=PatternFill('solid', fgColor='FFE699')))
    lm.freeze_panes = 'B6'

    # ---------------- Question Guide ----------------
    qg = wb.create_sheet('Question Guide')
    title(qg, 'Question Guide',
          'How to assess each question: what to ask, what evidence proves the level, what to watch for and how to reach the next level.')
    GH = ['ID', 'Service', 'Dimension', 'Question', 'Probing questions', 'Evidence and how to obtain it', 'Red flags',
          'Healthcare considerations', 'Path to 2 - Developing', 'Path to 3 - Defined', 'Path to 4 - Optimized',
          'References', 'Sources']
    GW = [8, 16, 18, 36, 52, 48, 36, 48, 30, 30, 30, 26, 40]
    header(qg, 4, GH, GW)
    qg.freeze_panes = 'E5'
    for r, q in enumerate(questions, 5):
        vals = [q['id'], q['service'], q['dimension'], q['question'], bullets(q['probing_questions']), bullets(q['evidence']),
                bullets(q['red_flags']), bullets(q['healthcare_considerations']), q['action_to_2'], q['action_to_3'],
                q['action_to_4'], q['references'], '\n'.join(split_list(q['sources']))]
        for c, v in enumerate(vals, 1):
            cell = qg.cell(row=r, column=c, value=v)
            cell.font = font(bold=(c == 1))
            cell.alignment = CENTER if c == 1 else WRAP
            cell.border = BORDER
        fit_height(qg, r, list(zip(vals, GW)))
    qg.print_title_rows = '4:4'

    # order: Start Here, Profile, LO, AM, ZT, Dashboard, Priority Gaps, CPG Crosswalk, Triage, License map, Guide, AllItems, Lists
    order = ['Start Here', 'Profile'] + list(SERVICES) + ['Dashboard', 'Priority Gaps', 'CPG Crosswalk',
                                                          'App Portfolio Triage', 'License-Security Map', 'Question Guide',
                                                          'AllItems', 'Lists']
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.title = 'ACS Health Maturity Kit — Assessment Workbook'
    wb.properties.creator = 'Andrade Consulting Services'
    wb.properties.version = VERSION
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'Wrote {OUT} ({len(questions)} questions)')


if __name__ == '__main__':
    build()
