#!/usr/bin/env python3
"""Score a completed ACS Health Maturity Assessment workbook.

Reads the answers directly (no Excel recalculation needed) and writes:
  - <output>.json : machine-readable results for the agents
  - <output>.md   : a readable summary

Usage:  python tools/score_assessment.py <completed.xlsx> [output_basename]
"""
import json, sys
from collections import OrderedDict, defaultdict
from pathlib import Path
from openpyxl import load_workbook

SERVICES = ['License Optimization', 'App Modernization', 'Zero Trust']
LEVELS = ['Not in place', 'Initial', 'Developing', 'Defined', 'Optimized']
CPGS_ESSENTIAL = ['Mitigate Known Vulnerabilities', 'Email Security', 'Multifactor Authentication',
                  'Basic Cybersecurity Training', 'Strong Encryption', 'Revoke Credentials for Departing Workforce Members',
                  'Basic Incident Planning and Preparedness', 'Unique Credentials', 'Separate User and Privileged Accounts',
                  'Vendor/Supplier Cybersecurity Requirements']
CPGS_ENHANCED = ['Asset Inventory', 'Third Party Vulnerability Disclosure', 'Third Party Incident Reporting',
                 'Cybersecurity Testing', 'Cybersecurity Mitigation', 'Detect and Respond to Relevant Threats and TTPs',
                 'Network Segmentation', 'Centralized Log Collection', 'Centralized Incident Planning and Preparedness',
                 'Configuration Management']


def level(score):
    return 'Not assessed' if score is None else LEVELS[min(4, int(score))]


def wavg(items):
    den = sum(i['weight'] for i in items if i['score'] is not None)
    return None if den == 0 else sum(i['score'] * i['weight'] for i in items if i['score'] is not None) / den


def read(path):
    wb = load_workbook(path, data_only=False)
    items = []
    for svc in SERVICES:
        ws = wb[svc]
        for r in range(5, ws.max_row + 1):
            qid = ws.cell(r, 1).value
            if not qid:
                continue
            ans = ws.cell(r, 9).value
            score = None
            if isinstance(ans, str) and ans[:1].isdigit():
                score = int(ans[0])
            elif isinstance(ans, (int, float)):
                score = int(ans)
            items.append(dict(service=svc, id=qid, dimension=ws.cell(r, 2).value, question=ws.cell(r, 3).value,
                              target_state=ws.cell(r, 4).value, hhs_cpg=ws.cell(r, 6).value or '',
                              weight=int(ws.cell(r, 8).value), answer=ans, notes=ws.cell(r, 10).value,
                              recommended_action=ws.cell(r, 11).value, score=score))
    profile = {}
    pf = wb['Profile']
    for r in range(4, 15):
        if pf.cell(r, 1).value:
            v = pf.cell(r, 2).value
            profile[pf.cell(r, 1).value] = v.isoformat() if hasattr(v, 'isoformat') else v
    lic = []
    lm = wb['License-Security Map']
    for r in range(6, lm.max_row + 1):
        cap = lm.cell(r, 1).value
        if not cap:
            continue
        licensed, deployed = lm.cell(r, 5).value, lm.cell(r, 6).value
        if licensed == 'Yes' and deployed in ('Partial', 'No'):
            status = 'Licensed but not fully deployed (quick win)'
        elif licensed == 'Yes' and deployed == 'Yes':
            status = 'Deployed'
        elif licensed == 'No':
            status = 'Not licensed'
        elif licensed:
            status = 'Confirm licensing'
        else:
            status = 'Not answered'
        lic.append(dict(capability=cap, licensing=lm.cell(r, 2).value, licensed=licensed, deployed=deployed, status=status))
    apps = []
    tp = wb['App Portfolio Triage']
    for r in range(6, tp.max_row + 1):
        name = tp.cell(r, 1).value
        if not name or str(name).startswith('EXAMPLE'):
            continue
        g = lambda c: tp.cell(r, c).value
        risk = {'No': 3, 'Unknown': 2}.get(g(6), 0) + {'Yes': 3, 'Unknown': 2}.get(g(7), 0) + (1 if g(8) == 'Yes' else 0) + (1 if (g(9) or 0) > 5 else 0)
        crit = g(10)
        apps.append(dict(application=name, hosting=g(4), platform=g(5), technical_risk=risk, criticality=crit,
                         business_value=g(11), priority=(risk * crit) if crit else None, final_decision=g(16)))
    return items, profile, lic, apps


def summarize(items, profile, lic, apps):
    res = OrderedDict(profile=profile, services=OrderedDict(), cpg=OrderedDict(), top_gaps=[], license_map=lic,
                      top_applications=sorted([a for a in apps if a['priority'] is not None], key=lambda a: -a['priority'])[:10])
    for svc in SERVICES:
        s_items = [i for i in items if i['service'] == svc]
        dims = OrderedDict()
        for d in OrderedDict.fromkeys(i['dimension'] for i in s_items):
            sc = wavg([i for i in s_items if i['dimension'] == d])
            dims[d] = dict(score=None if sc is None else round(sc, 2), level=level(sc))
        sc = wavg(s_items)
        res['services'][svc] = dict(score=None if sc is None else round(sc, 2), level=level(sc),
                                    answered=sum(i['score'] is not None for i in s_items), questions=len(s_items),
                                    dimensions=dims)
    sc = wavg(items)
    res['overall'] = dict(score=None if sc is None else round(sc, 2), level=level(sc),
                          answered=sum(i['score'] is not None for i in items), questions=len(items))
    for kind, goals in (('Essential', CPGS_ESSENTIAL), ('Enhanced', CPGS_ENHANCED)):
        for g in goals:
            tag = f'CPG-{"E" if kind == "Essential" else "X"}: {g}'
            mapped = [i for i in items if tag in i['hhs_cpg']]
            scores = [i['score'] for i in mapped if i['score'] is not None]
            avg = sum(scores) / len(scores) if scores else None
            status = 'Not assessed' if avg is None else 'Meets' if avg >= 3 else 'Partially meets' if avg >= 2 else 'Gap'
            res['cpg'][f'{kind}: {g}'] = dict(questions=[i['id'] for i in mapped],
                                              average=None if avg is None else round(avg, 2), status=status)
    gaps = [dict(id=i['id'], service=i['service'], dimension=i['dimension'], question=i['question'], score=i['score'],
                 weight=i['weight'], gap_points=(4 - i['score']) * i['weight'], recommended_action=i['recommended_action'],
                 hhs_cpg=i['hhs_cpg'], notes=i['notes']) for i in items if i['score'] is not None and i['score'] < 4]
    res['top_gaps'] = sorted(gaps, key=lambda g: -g['gap_points'])[:20]
    res['unanswered'] = [i['id'] for i in items if i['score'] is None and i['answer'] != 'N/A']
    return res


def to_markdown(res):
    L = ['# ACS Health Maturity Assessment — Results', '']
    org = res['profile'].get('Organization name') or 'Organization'
    L += [f'**Organization:** {org}  ', f'**Assessment date:** {res["profile"].get("Assessment date") or "—"}', '']
    o = res['overall']
    L += [f'**Overall maturity:** {o["score"]} / 4 ({o["level"]}) — {o["answered"]} of {o["questions"]} questions answered', '']
    L += ['## Scores by service', '', '| Service | Score (0–4) | Level | Answered |', '|---|---|---|---|']
    for s, v in res['services'].items():
        L.append(f'| {s} | {v["score"]} | {v["level"]} | {v["answered"]}/{v["questions"]} |')
    for s, v in res['services'].items():
        L += ['', f'### {s} by dimension', '', '| Dimension | Score | Level |', '|---|---|---|']
        for d, dv in v['dimensions'].items():
            L.append(f'| {d} | {dv["score"]} | {dv["level"]} |')
    L += ['', '## HHS Cybersecurity Performance Goals', '', '| Goal | Average | Status | Questions |', '|---|---|---|---|']
    for g, v in res['cpg'].items():
        L.append(f'| {g} | {v["average"]} | {v["status"]} | {", ".join(v["questions"])} |')
    L += ['', '## Top gaps', '', '| # | ID | Service | Score | Gap points | Recommended action |', '|---|---|---|---|---|---|']
    for n, g in enumerate(res['top_gaps'], 1):
        L.append(f'| {n} | {g["id"]} | {g["service"]} | {g["score"]} | {g["gap_points"]} | {g["recommended_action"]} |')
    quick = [l for l in res['license_map'] if l['status'].startswith('Licensed but')]
    if quick:
        L += ['', '## License quick wins (licensed but not fully deployed)', '']
        L += [f'- {l["capability"]} — {l["licensing"]}' for l in quick]
    if res['top_applications']:
        L += ['', '## Highest-priority applications', '', '| Application | Technical risk | Criticality | Priority |', '|---|---|---|---|']
        for a in res['top_applications']:
            L.append(f'| {a["application"]} | {a["technical_risk"]} | {a["criticality"]} | {a["priority"]} |')
    if res['unanswered']:
        L += ['', f'**Unanswered questions:** {", ".join(res["unanswered"])}']
    return '\n'.join(L) + '\n'


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    base = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix('')
    res = summarize(*read(src))
    Path(f'{base}.json').write_text(json.dumps(res, indent=2, default=str), encoding='utf-8')
    Path(f'{base}.md').write_text(to_markdown(res), encoding='utf-8')
    print(f'Wrote {base}.json and {base}.md — overall {res["overall"]["score"]} ({res["overall"]["level"]})')


if __name__ == '__main__':
    main()
